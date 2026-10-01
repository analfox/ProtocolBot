"""
mintrud.py - протокол проверки знаний по охране труда (Минтруд, Программа В)
и реестр обученных лиц для загрузки в Минтруд.

Протокол делается по шаблону mintrud_protocol_template.docx (вёрстка не
трогается). Плейсхолдеры:
  шапка:   {protocol_number} {day} {month} {year} {hours} {study_group}
  таблица: {name} {position} {workplace} {group}
Строка таблицы с {name} копируется на каждого участника; нумерация - списком
Word. Рег. номер из реестра не заполняется.

Реестр заполняется по шаблону registry_template.xlsx (копия Шаблон_Реестр_ОТ
из MintrudBot) начиная с 5-й строки.
"""
import copy
import os
import re

from docx import Document
from docx.oxml.ns import qn

from certificates import _fill_element, _date_parts, _strip_ids, suppress_hyphens_for

PROTOCOL_TEMPLATE = "mintrud_protocol_template.docx"
REGISTRY_TEMPLATE = "registry_template.xlsx"

REGISTRY_PROGRAM = "9. Безопасные методы и приемы выполнения работ на высоте"
REGISTRY_RESULT = "Удовлетворительно"
REGISTRY_FIRST_ROW = 5


def _template_path(name):
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), name)
    if not os.path.isfile(path):
        raise FileNotFoundError(f"Рядом с программой нет шаблона {name}")
    return path


def hours_text(hours):
    """16 -> '16 часов', 24 -> '24 часа', 21 -> '21 час'."""
    h = str(hours or "").strip()
    if not h.isdigit():
        return f"{h} часов" if h else "часов"
    n = int(h)
    if n % 10 == 1 and n % 100 != 11:
        word = "час"
    elif 2 <= n % 10 <= 4 and not 12 <= n % 100 <= 14:
        word = "часа"
    else:
        word = "часов"
    return f"{n} {word}"


def format_snils(raw):
    """Любая запись СНИЛС -> 'XXX-XXX-XXX XX'. Если цифр не 11 - как есть.
    Из Excel СНИЛС числом теряет ведущие нули (043... -> 43...), их возвращаем."""
    s = str(raw or "").strip()
    if re.fullmatch(r"\d+\.0", s):
        s = s[:-2]
    digits = re.sub(r"\D", "", s)
    if digits and 9 <= len(digits) < 11 and re.fullmatch(r"\d+", s):
        digits = digits.zfill(11)
    if len(digits) != 11:
        return s
    return f"{digits[:3]}-{digits[3:6]}-{digits[6:9]} {digits[9:]}"


def workplace(participant, organization=""):
    """Место работы: подразделение, а если его нет - организация из заявки."""
    return (participant.get("subdivision") or "").strip() or (organization or "").strip()


def create_mintrud_protocol(output_path, rows, info):
    """rows: [(участник, [группы])]; info: protocol_number, date, hours, study_group, organization."""
    doc = Document(_template_path(PROTOCOL_TEMPLATE))

    parts = _date_parts(info.get("date"))
    if parts:
        day, month_name, year = parts
        day = f"{day:02d}"
    else:
        day, month_name, year = info.get("date", ""), "", ""

    header = {
        "protocol_number": str(info.get("protocol_number", "")).strip(),
        "day": day,
        "month": month_name,
        "year": str(year),
        "hours": hours_text(info.get("hours")),
        "study_group": str(info.get("study_group", "")).strip(),
    }
    for p in doc.paragraphs:
        _fill_element(p._p, header)

    table = doc.tables[0]
    template_tr = next(
        (r._tr for r in table.rows if "{name}" in "".join(t.text or "" for t in r._tr.iter(qn("w:t")))),
        None,
    )
    if template_tr is None:
        raise ValueError(f"В шаблоне {PROTOCOL_TEMPLATE} нет строки таблицы с {{name}}")
    suppress_hyphens_for(template_tr, ("name",))

    org = info.get("organization", "")
    for i, (participant, groups) in enumerate(rows):
        tr = copy.deepcopy(template_tr)
        _strip_ids(tr, i)
        _fill_element(tr, {
            "name": participant.get("name", ""),
            "position": participant.get("position", ""),
            "workplace": workplace(participant, org),
            "group": ", ".join(groups),
        })
        template_tr.addprevious(tr)
    template_tr.getparent().remove(template_tr)

    doc.save(output_path)


def create_registry(output_path, participants, info):
    """Заполняет реестр ОТ для Минтруда: по строке на человека."""
    import openpyxl

    wb = openpyxl.load_workbook(_template_path(REGISTRY_TEMPLATE))
    ws = wb["Реестр"]

    org = info.get("organization", "")
    inn = str(info.get("inn", "")).strip()
    date = str(info.get("date", "")).strip()
    protocol = str(info.get("protocol_number", "")).strip()

    for i, p in enumerate(participants):
        r = REGISTRY_FIRST_ROW + i
        ws.cell(r, 1, (org or "").strip() or workplace(p))
        ws.cell(r, 2, inn)
        ws.cell(r, 3, p.get("name", ""))
        ws.cell(r, 4, format_snils(p.get("snils", "")))
        ws.cell(r, 5, p.get("position", ""))
        ws.cell(r, 6, REGISTRY_PROGRAM)
        ws.cell(r, 7, REGISTRY_RESULT)
        ws.cell(r, 8, date)
        ws.cell(r, 9, protocol)

    # openpyxl не сохраняет значения формул - пусть Excel пересчитает проверки при открытии.
    wb.calculation.fullCalcOnLoad = True
    wb.save(output_path)
